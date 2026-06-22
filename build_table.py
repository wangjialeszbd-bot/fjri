import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

# Each row: name, address, city/zip, beds, rent range (USD/mo), distance miles, distance km, phone
rows = [
    ["Eaves Mission Viejo", "24950 Via Florecer", "Mission Viejo, CA 92692", "1–2 BR", "$2,190 – $3,110", 1.23, 1.98, "(949) 525-4714"],
    ["Alicia Viejo Apartment Homes", "23842 Alicia Pkwy", "Mission Viejo, CA 92691", "1–2 BR", "$2,405 – $3,035", 1.94, 3.12, "(949) 472-8552"],
    ["Laurel Canyon Apartment Homes", "76 Mercantile Way", "Ladera Ranch, CA 92694", "1–3 BR", "$2,700 – $4,165", 2.06, 3.32, "(949) 485-3021"],
    ["Vista Del Lago", "21622 Marguerite Pkwy", "Mission Viejo, CA 92692", "1–3 BR", "$2,525 – $4,382", 2.45, 3.94, "(949) 243-7570"],
    ["Park Ridge Villas / Mosaic Apt Homes", "27444 Camden", "Mission Viejo, CA 92692", "Studio–2 BR", "$2,145 – $3,120", 2.79, 4.49, "(949) 393-5828"],
    ["Laurel Vista Apartment Homes", "27082 O'Neill Dr", "Ladera Ranch, CA 92694", "1–2 BR", "from $2,590", 2.99, 4.81, "(949) 682-7961"],
    ["Los Alisos at Mission Viejo", "28601 Los Alisos Blvd", "Mission Viejo, CA 92692", "Studio–3 BR", "$2,221 – $3,239", 3.02, 4.86, "(949) 762-8930"],
    ["Forest Glen (Lyon Forest Glen)", "25092 Farthing St", "Lake Forest, CA 92630", "1–2 BR", "$2,033 – $3,169", 3.03, 4.88, "(949) 393-5848"],
    ["Ridgecrest", "21486 Lake Forest Dr", "Lake Forest, CA 92630", "1–2 BR", "$2,507 – $2,945", 3.11, 5.00, "(949) 393-5814"],
]

headers = ["社区/房源名称\nProperty", "地址\nAddress", "城市/邮编\nCity & ZIP",
           "户型\nBeds", "月租范围\nMonthly Rent", "距离(英里)\nMiles",
           "距离(公里)\nKm", "电话\nPhone"]

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Saddleback 5km Rentals"

# Title
ws.merge_cells("A1:H1")
t = ws["A1"]
t.value = "Saddleback College 周边 5 公里租房房源  (Saddleback College, 28000 Marguerite Pkwy, Mission Viejo, CA 92692)"
t.font = Font(bold=True, size=13, color="FFFFFF")
t.fill = PatternFill("solid", fgColor="8B0000")
t.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
ws.row_dimensions[1].height = 38

ws.merge_cells("A2:H2")
s = ws["A2"]
s.value = "截至 2026-06。数据来自 Saddleback College 官方公寓列表 + 各租房网站公开报价；租金为大致区间，会随空置情况变动，签约前请电话确认。"
s.font = Font(italic=True, size=9, color="555555")
s.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
ws.row_dimensions[2].height = 26

# Header row
hr = 3
thin = Side(style="thin", color="BBBBBB")
border = Border(left=thin, right=thin, top=thin, bottom=thin)
for c, h in enumerate(headers, start=1):
    cell = ws.cell(row=hr, column=c, value=h)
    cell.font = Font(bold=True, color="FFFFFF", size=10)
    cell.fill = PatternFill("solid", fgColor="C0392B")
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell.border = border
ws.row_dimensions[hr].height = 34

# Data rows
for i, r in enumerate(rows):
    rr = hr + 1 + i
    for c, v in enumerate(r, start=1):
        cell = ws.cell(row=rr, column=c, value=v)
        cell.border = border
        cell.alignment = Alignment(horizontal="center" if c in (4,6,7) else "left",
                                   vertical="center", wrap_text=True)
        cell.font = Font(size=10)
    if i % 2 == 1:
        for c in range(1, 9):
            ws.cell(row=rr, column=c).fill = PatternFill("solid", fgColor="FBEEE9")
    ws.row_dimensions[rr].height = 22

widths = [34, 22, 24, 13, 18, 12, 12, 16]
for i, w in enumerate(widths, start=1):
    ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w

ws.freeze_panes = "A4"

out = "/home/user/fjri/output/Saddleback_College_5km_Rentals.xlsx"
wb.save(out)
print("saved", out)
